"""Lookup and calculation for source-backed cooking yield factors."""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook


@dataclass(frozen=True)
class CookingFactor:
    food_name: str
    cooking_method: str
    raw_form: str
    cooked_form: str
    yield_factor: float
    data_status: str
    source_reference: str
    confidence: str


@dataclass(frozen=True)
class FoodQuantityCalculation:
    prescribed_quantity: float
    calculated_quantity: float
    weight_state: Optional[str] = None
    cooking_factor: Optional[CookingFactor] = None


_WORKBOOK_PATH = Path(__file__).resolve().parents[4] / "docs" / "cooking_factors_dieta.xlsx"
_PREPARATION_WORDS = {
    "assado": ("prepared", None),
    "assada": ("prepared", None),
    "cozido": ("prepared", "boiling"),
    "cozida": ("prepared", "boiling"),
    "grelhado": ("prepared", "grilling"),
    "grelhada": ("prepared", "grilling"),
    "refogado": ("prepared", "sauté"),
    "refogada": ("prepared", "sauté"),
    "cru": ("raw", "raw"),
    "crua": ("raw", "raw"),
}


def _normalize(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value.casefold())
    without_marks = "".join(char for char in decomposed if not unicodedata.combining(char))
    return " ".join(re.findall(r"[a-z0-9]+", without_marks))


def _food_identity(name: str) -> tuple[str, Optional[str], Optional[str]]:
    words = _normalize(name).split()
    state = None
    method = None
    while words and words[-1] in _PREPARATION_WORDS:
        state, method = _PREPARATION_WORDS[words.pop()]
    return " ".join(words), state, method


@lru_cache(maxsize=1)
def _load_factors() -> tuple[CookingFactor, ...]:
    if not _WORKBOOK_PATH.is_file():
        return ()

    workbook = load_workbook(_WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        sheet = workbook["CookingFactor"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        factors = []
        for values in rows:
            record = dict(zip(headers, values))
            required = (
                "food_name", "cooking_method", "raw_form", "cooked_form",
                "yield_factor", "data_status", "source_reference", "confidence",
            )
            if any(record.get(field) is None for field in required):
                continue
            factors.append(CookingFactor(
                food_name=str(record["food_name"]),
                cooking_method=str(record["cooking_method"]),
                raw_form=str(record["raw_form"]),
                cooked_form=str(record["cooked_form"]),
                yield_factor=float(record["yield_factor"]),
                data_status=str(record["data_status"]),
                source_reference=str(record["source_reference"]),
                confidence=str(record["confidence"]),
            ))
        return tuple(factors)
    finally:
        workbook.close()


def find_cooking_factor(
    food: str,
    cooking_method: Optional[str],
    prescribed_state: str,
    raw_form: Optional[str] = None,
    cooked_form: Optional[str] = None,
    allow_unique_method_fallback: bool = False,
) -> Optional[CookingFactor]:
    """Return only a unique compatible row; never guess between methods."""
    food_key, inferred_state, inferred_method = _food_identity(food)
    method = cooking_method or inferred_method
    state = prescribed_state or inferred_state
    if state not in {"raw", "prepared"}:
        return None

    candidates = []
    for factor in _load_factors():
        if factor.data_status == "not_applicable":
            continue
        factor_food = _normalize(factor.food_name)
        name_matches = (
            factor_food == food_key
            or factor_food.startswith(food_key + " ")
            or food_key.startswith(factor_food + " ")
        )
        if factor.yield_factor <= 0 or not name_matches:
            continue
        if raw_form and _normalize(factor.raw_form) != _normalize(raw_form):
            continue
        if cooked_form and _normalize(factor.cooked_form) != _normalize(cooked_form):
            continue
        candidates.append(factor)

    if method:
        method_matches = [
            factor for factor in candidates
            if (
                _normalize(factor.cooking_method) == _normalize(method)
                or _normalize(method) in {
                    _normalize(part) for part in factor.cooking_method.split("_")
                }
            )
        ]
        if len(method_matches) == 1:
            return method_matches[0]
        if method_matches or not allow_unique_method_fallback:
            return None
    return candidates[0] if len(candidates) == 1 else None


def calculate_food_quantity(
    prescribed_quantity: float,
    prescribed_state: Optional[str],
    target_state: str,
    food: str,
    cooking_method: Optional[str] = None,
    consider_cooking_factor: bool = False,
    raw_form: Optional[str] = None,
    cooked_form: Optional[str] = None,
) -> FoodQuantityCalculation:
    """Calculate a target weight without changing the prescribed quantity."""
    inferred_food, inferred_state, inferred_method = _food_identity(food)
    del inferred_food
    state = prescribed_state or inferred_state
    method = cooking_method or inferred_method
    if not consider_cooking_factor or state not in {"raw", "prepared"} or state == target_state:
        if not consider_cooking_factor or state == "raw" or state == target_state:
            return FoodQuantityCalculation(prescribed_quantity, prescribed_quantity, state)
        if state is None:
            state = "prepared"

    factor = find_cooking_factor(
        food,
        method,
        state,
        raw_form=raw_form,
        cooked_form=cooked_form,
        allow_unique_method_fallback=(cooking_method is None and method == "boiling"),
    )
    if factor is None:
        return FoodQuantityCalculation(prescribed_quantity, prescribed_quantity, state)

    if state == "prepared" and target_state == "raw":
        calculated = prescribed_quantity / factor.yield_factor
    elif state == "raw" and target_state == "prepared":
        calculated = prescribed_quantity * factor.yield_factor
    else:
        calculated = prescribed_quantity
    return FoodQuantityCalculation(prescribed_quantity, calculated, state, factor)